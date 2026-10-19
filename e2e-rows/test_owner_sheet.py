"""CLI-only owner-sheet rows. Run on the disposable Linux box, never the Mac.

All inputs, product bytes, CLI outputs and exits are retained in the printed
evidence directory. --mutant-drop-answers changes only a disposable copy of
the CLI; the preservation row must then fail at its answer-preservation check.
"""

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import uuid

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as error:
    if error.name != "openpyxl":
        raise
    load_workbook = None


MUTANT = False
SOURCE = Path(__file__).resolve().parents[1] / "integrations/owner-sheet/owner_sheet.py"
EN_COLUMNS = (
    "ID", "Priority", "Category", "Question", "My recommendation", "Your answer",
    "Note (optional)", "What this blocks", "Default if unanswered", "Where / time needed", "Status",
)
ZH_COLUMNS = (
    "編號", "優先", "類別", "事項", "我的建議", "你的回答", "備註（可留空）",
    "會卡住什麼", "你沒回答時我的做法", "在哪裡做／需時", "狀態",
)


def setUpModule():
    if sys.platform != "linux":
        raise RuntimeError("owner-sheet runtime rows require the disposable Linux test box")
    if load_workbook is None:
        message = "SKIP OWNER-SHEET ROWS: openpyxl is absent; install it on the box before claiming a pass"
        print(message, file=sys.stderr, flush=True)
        raise unittest.SkipTest(message)


class OwnerSheetRows(unittest.TestCase):
    def setUp(self):
        evidence = os.environ.get("OWNER_SHEET_EVIDENCE")
        if evidence:
            parent = Path(evidence)
            self.assertTrue(parent.is_absolute(), "OWNER_SHEET_EVIDENCE must be absolute")
            parent.mkdir(parents=True, exist_ok=True)
        else:
            parent = Path(tempfile.gettempdir())
        self.root = Path(tempfile.mkdtemp(prefix="lb-plug-sheet-row-", dir=parent))
        print("owner-sheet evidence retained: " + str(self.root), file=sys.stderr, flush=True)
        (self.root / "row.txt").write_text(self._testMethodName + "\n", encoding="utf-8")
        source = SOURCE.read_bytes()
        if MUTANT:
            original = b"row.values[5:7] = previous.values[5:7]"
            self.assertEqual(source.count(original), 1, "mutant replacement must hit exactly once")
            source = source.replace(original, b'row.values[5:7] = ["", ""]')
        self.script = self.root / "owner_sheet.py"
        self.script.write_bytes(source)
        (self.root / "source.sha256").write_text(hashlib.sha256(source).hexdigest() + "\n")
        self.home = self.root / "home"
        self.home.mkdir()
        self.environment = {"HOME": str(self.home), "LANG": "C.UTF-8", "PYTHONIOENCODING": "utf-8"}
        self.config = self.root / "config.json"
        self.queue = self.root / "queue.json"
        self.sheet = self.root / "synced folder/owner.xlsx"
        self.state = self.root / "state.json"
        self.settings = {"queue": "queue.json", "sheet": "synced folder/owner.xlsx", "state": "state.json"}
        self.write_config()
        self.ids = ["00" + uuid.uuid4().hex[:10] for _ in range(3)]
        self.questions = [self.question(identifier) for identifier in self.ids]
        self.write_queue(self.questions)
        self.calls = 0

    def question(self, identifier):
        return {
            "id": identifier, "priority": "high", "category": "planning",
            "question": "Which option? " + uuid.uuid4().hex[:8], "recommendation": "a",
            "blocks": "next step", "default_if_unanswered": "wait", "where": "this sheet",
        }

    def write_config(self):
        self.config.write_text(json.dumps(self.settings, ensure_ascii=False), encoding="utf-8")

    def write_queue(self, questions):
        self.queue.write_text(json.dumps({"rows": questions}, ensure_ascii=False), encoding="utf-8")

    def cli(self, *arguments, isolated=False, config=None, cwd=None):
        self.calls += 1
        record = self.root / ("call-" + str(self.calls))
        record.mkdir()
        selected_config = config or self.config
        argv = [sys.executable, *(["-I", "-S"] if isolated else []), str(self.script),
                "--config", str(selected_config), *map(str, arguments)]
        (record / "argv.json").write_text(json.dumps(argv) + "\n", encoding="utf-8")
        for name, path in (("config", selected_config), ("queue", self.queue), ("sheet", self.sheet), ("state", self.state)):
            if path.exists():
                (record / (name + ".input")).write_bytes(path.read_bytes())
        result = subprocess.run(argv, cwd=cwd or self.root, env=self.environment,
                                capture_output=True, timeout=30)
        (record / "stdout").write_bytes(result.stdout)
        (record / "stderr").write_bytes(result.stderr)
        (record / "exit").write_text(str(result.returncode) + "\n")
        result.stdout = result.stdout.decode("utf-8")
        result.stderr = result.stderr.decode("utf-8")
        return result

    def ok(self, *arguments, **options):
        result = self.cli(*arguments, **options)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stderr, "")
        return result.stdout

    def refused(self, result, reason):
        self.assertEqual(result.returncode, 2, result.stderr)
        self.assertIn(reason, result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertEqual(result.stdout, "", "refused command must not emit partial rows or success")

    def read(self, *identifiers):
        return [json.loads(line) for line in self.ok("read", *identifiers).splitlines()]

    def answer(self, entries, worksheet="Questions"):
        workbook = load_workbook(self.sheet)
        try:
            matched = set()
            for cells in workbook[worksheet].iter_rows(min_row=2):
                if cells[0].value in entries:
                    cells[5].value, cells[6].value = entries[cells[0].value]
                    matched.add(cells[0].value)
            self.assertEqual(matched, set(entries), "owner edit could not find every requested row")
            workbook.save(self.sheet)
        finally:
            workbook.close()

    def rows(self, worksheet="Questions"):
        workbook = load_workbook(self.sheet, data_only=False)
        try:
            return {cells[0].value: tuple(cell.value for cell in cells)
                    for cells in workbook[worksheet].iter_rows(min_row=2) if cells[0].value is not None}
        finally:
            workbook.close()

    def expected(self, identifier, answer="", note=""):
        return {"id": identifier, "answer": answer, "note": note}

    def test_read_ack_changes(self):
        self.ok("build")
        self.assertEqual(self.read(), [], "untouched rows must not appear")
        first, second, _ = self.ids
        self.answer({first: ("  OK\n", "keep spacing  "), second: (None, "note only")})
        self.assertEqual(self.read(), [self.expected(first, "  OK\n", "keep spacing  "),
                                       self.expected(second, note="note only")])
        self.assertEqual(self.read(second), [self.expected(second, note="note only")])
        self.ok("ack", first, second)
        self.assertEqual(self.read(), [], "acknowledged rows must disappear")
        self.answer({first: ("  OK\n", "changed note")})
        self.assertEqual(self.read(), [self.expected(first, "  OK\n", "changed note")])
        self.ok("ack", first)
        self.assertEqual(self.read(), [])
        self.answer({first: (None, None)})
        self.assertEqual(self.read(), [self.expected(first)], "clearing an acknowledged answer must be reported")
        self.ok("ack", first)
        self.assertEqual(self.read(), [])
        self.answer({second: ("b", "note only")})
        self.assertEqual(self.read(), [self.expected(second, "b", "note only")], "answer-only changes must be reported")

    def test_rebuild_preserves_answers_by_id(self):
        self.questions[0]["question"] = "=1+1"
        self.ok("build")
        pairs = {self.ids[0]: ("  OK\n", "不要刪除  "), self.ids[1]: (0, False), self.ids[2]: ("=1+2", "c")}
        self.answer(pairs)
        workbook = load_workbook(self.sheet)
        try:
            for cells in workbook["Questions"].iter_rows(min_row=2):
                if cells[0].value == self.ids[1]:
                    cells[5].number_format = "0.00"
            self.assertEqual(workbook["Questions"].cell(2, 4).data_type, "s", "queue text became an executable formula")
            workbook.save(self.sheet)
        finally:
            workbook.close()
        before = self.rows()
        self.ok("ack", *self.ids)
        self.assertEqual(self.read(), [])
        added = self.question("new-" + uuid.uuid4().hex[:8])
        revised = dict(self.questions[0], question="Revised question", priority="low")
        self.write_queue([self.questions[2], added, revised, self.questions[1]])
        self.ok("build")
        after = self.rows()
        self.assertEqual(list(after), [self.ids[2], added["id"], self.ids[0], self.ids[1]])
        for identifier in self.ids:
            self.assertEqual(after[identifier][5:7], before[identifier][5:7], "build lost an owner answer or note")
        self.assertEqual(after[self.ids[0]][3], "Revised question")
        self.assertEqual(after[added["id"]][5:7], (None, None))
        self.assertEqual(self.read(), [], "metadata changes must not create new owner answers")
        workbook = load_workbook(self.sheet, data_only=False)
        try:
            cells = {row[0].value: row for row in workbook["Questions"].iter_rows(min_row=2)}
            self.assertEqual(cells[self.ids[1]][5].data_type, "n")
            self.assertEqual(cells[self.ids[1]][5].number_format, "0.00")
            self.assertEqual(cells[self.ids[1]][6].data_type, "b")
            self.assertEqual(cells[self.ids[2]][5].data_type, "f")
        finally:
            workbook.close()
        self.assertEqual(list(self.sheet.parent.glob(".*.tmp")), [], "completed build left temporary files")

    def test_closed_answers_survive_and_reopen(self):
        self.ok("build")
        first, second, third = self.ids
        self.answer({first: ("a", "keep closed"), second: (None, "closed note only")})
        self.write_queue([self.questions[2]])
        self.ok("build")
        self.assertEqual(list(self.rows()), [third])
        self.assertEqual(self.rows("Closed")[first][5:7], ("a", "keep closed"))
        self.assertEqual(self.rows("Closed")[second][5:7], (None, "closed note only"))
        self.assertEqual(self.read(), [self.expected(first, "a", "keep closed"),
                                       self.expected(second, note="closed note only")])
        self.ok("ack", first, second)
        self.assertEqual(self.read(), [])
        self.answer({second: ("c", "edited while closed")}, worksheet="Closed")
        self.ok("build")
        self.assertEqual(self.rows("Closed")[second][5:7], ("c", "edited while closed"))
        self.write_queue([self.questions[1], self.questions[2]])
        self.ok("build")
        self.assertEqual(self.rows()[second][5:7], ("c", "edited while closed"))
        self.assertEqual(list(self.rows("Closed")), [first])
        self.assertEqual(self.read(), [self.expected(second, "c", "edited while closed")])

    def test_unknown_ids_and_bad_queue(self):
        self.ok("build")
        self.answer({self.ids[0]: ("OK", "retain")})
        self.ok("ack", self.ids[0])
        self.answer({self.ids[0]: ("OK", "new unhandled note")})
        sheet, state = self.sheet.read_bytes(), self.state.read_bytes()
        missing = "unknown-" + uuid.uuid4().hex[:8]
        for command in ("read", "ack"):
            self.refused(self.cli(command, self.ids[0], missing), "unknown id: " + missing)
            self.assertEqual(self.sheet.read_bytes(), sheet)
            self.assertEqual(self.state.read_bytes(), state)
        self.write_queue([self.questions[0], self.questions[0]])
        self.refused(self.cli("build"), "duplicate queue id")
        self.assertEqual(self.sheet.read_bytes(), sheet, "invalid queue replaced the workbook")
        self.write_queue([dict(self.questions[0], question="x" * 32768)])
        self.refused(self.cli("build"), "Excel cannot preserve")
        self.assertEqual(self.sheet.read_bytes(), sheet, "oversized input replaced the workbook")

    def test_both_languages_and_language_switch(self):
        self.ok("build")
        self.answer({self.ids[0]: ("OK", "保留")})
        for language, names, columns, help_fragment in (
            ("zh-TW", ("待你決定", "已結束", "說明"), ZH_COLUMNS, "直接寫你的指示"),
            ("en", ("Questions", "Closed", "Help"), EN_COLUMNS, "write your instructions"),
        ):
            self.settings["language"] = language
            self.settings["title"] = "Review queue / 決策"
            self.write_config()
            self.ok("build")
            workbook = load_workbook(self.sheet)
            try:
                self.assertEqual(tuple(workbook.sheetnames), names)
                self.assertEqual(tuple(cell.value for cell in workbook[names[0]][1]), columns)
                help_text = "\n".join(str(row[0]) for row in workbook[names[2]].values)
                for fragment in ("OK", "a / b / c", help_fragment, self.settings["title"]):
                    self.assertIn(fragment, help_text)
                self.assertEqual(workbook.properties.title, self.settings["title"])
            finally:
                workbook.close()
            self.assertEqual(self.rows(names[0])[self.ids[0]][5:7], ("OK", "保留"))

    def test_init_and_toml_or_json_config(self):
        directory = self.root / "new setup"
        config = directory / "owner-sheet.json"
        self.ok("init", "--language", "zh-TW", config=config)
        self.assertEqual(json.loads((directory / "owner-queue.json").read_text()), {"rows": []})
        self.ok("build", config=config)
        sheet = directory / "owner-decisions.xlsx"
        workbook = load_workbook(sheet)
        try:
            self.assertIn("待你決定", workbook.sheetnames)
        finally:
            workbook.close()
        original = config.read_bytes()
        self.refused(self.cli("init", config=config), "init requires unused")
        self.assertEqual(config.read_bytes(), original)
        toml_config = self.root / "config.toml"
        toml_config.write_text('queue = "queue.json"\nsheet = "synced folder/owner.xlsx"\nstate = "state.json"\n', encoding="utf-8")
        if sys.version_info >= (3, 11):
            self.ok("build", config=toml_config, cwd=self.home)
            self.assertEqual(list(self.rows()), self.ids, "relative paths must resolve beside config")
            initialized = self.root / "toml setup/owner-sheet.toml"
            self.ok("init", "--title", "Decisions / 決策", config=initialized)
            self.ok("build", config=initialized)
            workbook = load_workbook(initialized.parent / "owner-decisions.xlsx")
            try:
                self.assertEqual(workbook["Help"].cell(1, 1).value, "Decisions / 決策")
                self.assertEqual(tuple(cell.value for cell in workbook["Questions"][1]), EN_COLUMNS)
            finally:
                workbook.close()
        else:
            self.refused(self.cli("build", config=toml_config), "use a .json config")
            self.ok("build", cwd=self.home)
            self.assertEqual(list(self.rows()), self.ids)

    def test_bad_workbook_and_state(self):
        self.ok("build")
        self.answer({self.ids[0]: ("OK", "keep")})
        workbook = load_workbook(self.sheet)
        try:
            workbook["Questions"].cell(3, 1, self.ids[0])
            workbook.save(self.sheet)
        finally:
            workbook.close()
        original = self.sheet.read_bytes()
        self.refused(self.cli("build"), "duplicate sheet id")
        self.assertEqual(self.sheet.read_bytes(), original)
        workbook = load_workbook(self.sheet)
        try:
            workbook["Questions"].cell(3, 1, self.ids[1])
            workbook.save(self.sheet)
        finally:
            workbook.close()
        self.state.write_text('{"version": 1, "answers": []}', encoding="utf-8")
        original = self.state.read_bytes()
        self.refused(self.cli("ack", self.ids[0]), "invalid ack state")
        self.assertEqual(self.state.read_bytes(), original)

    def test_missing_openpyxl_refuses_clearly(self):
        result = self.cli("build", isolated=True)
        self.refused(result, "openpyxl is required")
        self.assertIn("pip install openpyxl", result.stderr)
        self.assertFalse(self.sheet.exists(), "missing dependency still created a workbook")
        self.assertFalse(self.state.exists(), "missing dependency still created ack state")


if __name__ == "__main__":
    if "--mutant-drop-answers" in sys.argv:
        MUTANT = True
        sys.argv.remove("--mutant-drop-answers")
        print("MUTANT: build drops existing answers in a disposable CLI copy; preservation must fail", file=sys.stderr)
    unittest.main(verbosity=2)
