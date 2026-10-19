#!/usr/bin/env python3
"""A local, sync-folder owner decision sheet. Python 3.9+, openpyxl only."""

import argparse
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from zipfile import BadZipFile

try:
    import tomllib
except ModuleNotFoundError:
    tomllib = None


FIELDS = (
    "id", "priority", "category", "question", "recommendation", "answer",
    "note", "blocks", "default_if_unanswered", "where", "status",
)
QUEUE_FIELDS = tuple(key for key in FIELDS if key not in ("answer", "note", "status"))
WIDTHS = (16, 12, 16, 70, 42, 36, 28, 34, 34, 30, 24)
LABELS = {
    "en": {
        "sheets": ("Questions", "Closed", "Help"),
        "title": "Owner decisions",
        "columns": (
            "ID", "Priority", "Category", "Question", "My recommendation",
            "Your answer", "Note (optional)", "What this blocks",
            "Default if unanswered", "Where / time needed", "Status",
        ),
        "status": ("Awaiting answer", "New answer", "Handled", "Closed"),
        "help": (
            "How to use: type in Your answer when you have time; you need not answer every row at once.",
            "To accept My recommendation, type OK. For options, type a / b / c. Otherwise write your instructions.",
            "The supervisor reads new answers on its configured tick and acknowledges them after handling them. Status refreshes on the next build.",
            "Edit this .xlsx file in place. Do not convert it to a separate online spreadsheet, rename it or move it; the supervisor reads this exact file.",
            "Fill only Your answer and Note (optional). The supervisor maintains the other columns and preserves your entries when rebuilding.",
            "Closed contains questions removed from the queue. Their answers and notes remain readable and editable.",
            "Legal, payment and account actions require explicit authorization. An unanswered row is never consent.",
        ),
    },
    "zh-TW": {
        "sheets": ("待你決定", "已結束", "說明"),
        "title": "待你決定",
        "columns": (
            "編號", "優先", "類別", "事項", "我的建議", "你的回答", "備註（可留空）",
            "會卡住什麼", "你沒回答時我的做法", "在哪裡做／需時", "狀態",
        ),
        "status": ("待回答", "有新回答", "已處理", "已結束"),
        "help": (
            "怎麼用：在「你的回答」欄直接打字即可，有空再回，不必一次回完。",
            "想照我的建議：填「OK」。有選項的填 a / b / c。其他情況直接寫你的指示。",
            "主管會依設定的週期讀取新回答，處理後再確認；「狀態」在下次更新檔案時刷新。",
            "請直接編輯這個 .xlsx 檔案，不要另存成新的線上試算表、不要改檔名或搬位置，否則主管讀不到。",
            "只需要填「你的回答」與「備註」兩欄；其他欄位由主管維護，更新時會保留你填的內容。",
            "「已結束」保留已從佇列移除的事項，回答與備註仍可讀取及編輯。",
            "不可逆的事（法律、金流、帳號）需要明確授權；未回答不代表同意。",
        ),
    },
}


class Refusal(ValueError):
    pass


def text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def checked_text(value, name, nonempty=False):
    if not isinstance(value, str) or (nonempty and not value.strip()):
        raise Refusal(name + " must be " + ("a nonempty string" if nonempty else "a string"))
    if len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise Refusal(name + " contains text Excel cannot preserve")
    return value


def read_bytes(path):
    try:
        return path.read_bytes()
    except FileNotFoundError:
        return None


def atomic_write(path, writer, previous):
    """Save beside the destination, refuse observed concurrent edits, then rename."""
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, name = tempfile.mkstemp(prefix="." + path.name + ".", suffix=".tmp", dir=path.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        writer(temporary)
        with temporary.open("rb") as stream:
            os.fsync(stream.fileno())
        if read_bytes(path) != previous:
            raise Refusal("destination changed while writing; no replacement made; finish syncing and retry")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def save_json(path, value, previous):
    data = (json.dumps(value, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    atomic_write(path, lambda temporary: temporary.write_bytes(data), previous)


def load_config(path):
    raw = path.read_bytes()
    if path.suffix.lower() == ".json":
        config = json.loads(raw)
    elif tomllib is not None:
        config = tomllib.loads(raw.decode("utf-8"))
    else:
        raise Refusal("TOML configuration needs Python 3.11+; use a .json config on Python 3.9/3.10")
    if not isinstance(config, dict) or set(config) - {"queue", "sheet", "state", "language", "title"}:
        raise Refusal("config must contain only queue, sheet, state, language and optional title")
    for key in ("queue", "sheet", "state"):
        value = checked_text(config.get(key), "config " + key, nonempty=True)
        candidate = Path(value).expanduser()
        config[key] = (path.parent / candidate).resolve()
    if len({path.resolve(), config["queue"], config["sheet"], config["state"]}) != 4:
        raise Refusal("config, queue, sheet and state must have distinct paths")
    if config["sheet"].suffix.lower() != ".xlsx":
        raise Refusal("sheet must have an .xlsx extension")
    language = config.setdefault("language", "en")
    if not isinstance(language, str) or language not in LABELS:
        raise Refusal("language must be en or zh-TW")
    config.setdefault("title", LABELS[language]["title"])
    checked_text(config["title"], "title", nonempty=True)
    return config


def init(path, language, title):
    config = {
        "queue": "owner-queue.json", "sheet": "owner-decisions.xlsx",
        "state": "owner-sheet.state.json", "language": language,
        "title": title or LABELS[language]["title"],
    }
    checked_text(config["title"], "title", nonempty=True)
    paths = [path.resolve()] + [(path.parent / config[key]).resolve() for key in ("queue", "sheet", "state")]
    if len(set(paths)) != 4 or any(item.exists() for item in paths):
        raise Refusal("init requires unused config, queue, sheet and state paths; existing files are preserved")
    if path.suffix.lower() != ".json" and tomllib is None:
        raise Refusal("TOML configuration needs Python 3.11+; use a .json config on Python 3.9/3.10")
    if path.suffix.lower() == ".json":
        contents = json.dumps(config, ensure_ascii=False, indent=2) + "\n"
    else:
        contents = "".join(key + " = " + json.dumps(value, ensure_ascii=False) + "\n" for key, value in config.items())
    save_json(paths[1], {"rows": []}, None)
    atomic_write(path, lambda temporary: temporary.write_text(contents, encoding="utf-8"), None)
    print(json.dumps({"initialized": True}))


@dataclass
class SheetRow:
    values: list
    owner_types: tuple = ("s", "s")
    owner_formats: tuple = ("General", "General")

    @property
    def id(self):
        return self.values[0]

    def answer(self):
        return [text(value) for value in self.values[5:7]]


def queue_rows(path):
    queue = json.loads(path.read_bytes())
    if not isinstance(queue, dict) or set(queue) != {"rows"} or not isinstance(queue["rows"], list):
        raise Refusal('queue must be an object with a "rows" array')
    rows = {}
    for record in queue["rows"]:
        if not isinstance(record, dict) or set(record) - set(QUEUE_FIELDS):
            raise Refusal("each queue row must contain only the documented question fields")
        checked_text(record.get("id"), "row id", nonempty=True)
        checked_text(record.get("question"), "row question", nonempty=True)
        values = [checked_text(record.get(key, ""), "row " + key) for key in FIELDS]
        if record["id"] in rows:
            raise Refusal("duplicate queue id: " + record["id"])
        rows[record["id"]] = SheetRow(values)
    return rows


def sheet_rows(path):
    previous = read_bytes(path)
    if previous is None:
        return {}, None
    rows = {}
    workbook = load_workbook(BytesIO(previous), data_only=False)
    try:
        layouts = [labels for labels in LABELS.values() if set(workbook.sheetnames) == set(labels["sheets"])]
        if not layouts:
            raise Refusal("worksheet names changed; restore the generated Questions, Closed and Help worksheets (or their translated names)")
        for labels in layouts:
            for name in labels["sheets"][:2]:
                worksheet = workbook[name]
                if tuple(cell.value for cell in worksheet[1]) != labels["columns"]:
                    raise Refusal("question worksheet headers changed; restore the generated columns before rebuilding")
                for cells in worksheet.iter_rows(min_row=2):
                    values = [cell.value for cell in cells]
                    if all(value is None for value in values):
                        continue
                    identifier = checked_text(values[0], "sheet row id", nonempty=True)
                    if identifier in rows:
                        raise Refusal("duplicate sheet id: " + identifier)
                    for value in values[5:7]:
                        checked_text(text(value), "owner answer or note")
                    rows[identifier] = SheetRow(
                        values, tuple(cell.data_type for cell in cells[5:7]),
                        tuple(cell.number_format for cell in cells[5:7]),
                    )
    finally:
        workbook.close()
    return rows, previous


def ack_state(path):
    previous = read_bytes(path)
    if previous is None:
        return {}, None
    state = json.loads(previous)
    if (not isinstance(state, dict) or set(state) != {"version", "answers"}
            or type(state["version"]) is not int or state["version"] != 1
            or not isinstance(state["answers"], dict)):
        raise Refusal("invalid ack state; expected version 1 with an answers object")
    for identifier, pair in state["answers"].items():
        if (not identifier.strip() or not isinstance(pair, list) or len(pair) != 2
                or not all(isinstance(value, str) for value in pair)):
            raise Refusal("invalid ack state answer")
    return state["answers"], previous


def changed(row, state):
    pair = row.answer()
    return pair != state[row.id] if row.id in state else any(pair)


def build(config):
    current = queue_rows(config["queue"])
    old, previous_sheet = sheet_rows(config["sheet"])
    state, _ = ack_state(config["state"])
    for identifier, row in current.items():
        if identifier in old:
            previous = old[identifier]
            row.values[5:7] = previous.values[5:7]
            row.owner_types = previous.owner_types
            row.owner_formats = previous.owner_formats
    closed = [row for identifier, row in old.items() if identifier not in current]
    labels = LABELS[config["language"]]
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = config["title"]
    try:
        for name, rows, is_closed in ((labels["sheets"][0], list(current.values()), False),
                                      (labels["sheets"][1], closed, True)):
            worksheet = workbook.create_sheet(name)
            worksheet.append(labels["columns"])
            for row_number, row in enumerate(rows, 2):
                status = labels["status"][1 if changed(row, state) else 2 if row.id in state else 0]
                row.values[10] = labels["status"][3] + " / " + status if is_closed else status
                for index, value in enumerate(row.values, 1):
                    cell = worksheet.cell(row_number, index)
                    cell.value = value if index in (6, 7) else text(value)
                    cell.data_type = row.owner_types[index - 6] if index in (6, 7) else "s"
                    if index in (6, 7):
                        cell.number_format = row.owner_formats[index - 6]
            for index, width in enumerate(WIDTHS, 1):
                worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = width
            for cells in worksheet.iter_rows():
                for cell in cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.row == 1:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="D9E1F2")
                    elif cell.column in (6, 7):
                        cell.fill = PatternFill("solid", fgColor="FFF2CC")
            worksheet.freeze_panes = "D2"
            worksheet.auto_filter.ref = worksheet.dimensions
        help_sheet = workbook.create_sheet(labels["sheets"][2])
        help_sheet.column_dimensions["A"].width = 120
        for row_number, line in enumerate((config["title"], *labels["help"]), 1):
            cell = help_sheet.cell(row_number, 1, line)
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        atomic_write(config["sheet"], workbook.save, previous_sheet)
    finally:
        workbook.close()
    print(json.dumps({"built": len(current), "closed": len(closed)}))


def read_or_ack(config, command, identifiers):
    rows, previous_sheet = sheet_rows(config["sheet"])
    if previous_sheet is None:
        raise Refusal("sheet does not exist; run build first")
    unknown = [identifier for identifier in identifiers if identifier not in rows]
    if unknown:
        raise Refusal("unknown id: " + ", ".join(unknown))
    state, previous_state = ack_state(config["state"])
    selected = list(dict.fromkeys(identifiers)) if identifiers else list(rows)
    if command == "read":
        for identifier in selected:
            row = rows[identifier]
            if changed(row, state):
                answer, note = row.answer()
                print(json.dumps({"id": identifier, "answer": answer, "note": note}, ensure_ascii=False))
    else:
        for identifier in selected:
            state[identifier] = rows[identifier].answer()
        if read_bytes(config["sheet"]) != previous_sheet:
            raise Refusal("sheet changed during ack; no acknowledgement written; read again")
        save_json(config["state"], {"version": 1, "answers": state}, previous_state)
        print(json.dumps({"acknowledged": selected}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True, type=Path, help="TOML (Python 3.11+) or JSON configuration")
    commands = parser.add_subparsers(dest="command", required=True)
    initializer = commands.add_parser("init", help="create a new config and empty queue; never overwrite")
    initializer.add_argument("--language", choices=tuple(LABELS), default="en")
    initializer.add_argument("--title")
    commands.add_parser("build", help="rebuild, retaining owner entries and closed questions")
    reader = commands.add_parser("read", help="print changed answers/notes as JSON lines")
    reader.add_argument("ids", metavar="ID", nargs="*")
    acknowledger = commands.add_parser("ack", help="acknowledge current answers/notes after handling them")
    acknowledger.add_argument("ids", metavar="ID", nargs="+")
    arguments = parser.parse_args()
    try:
        global Workbook, load_workbook, Alignment, Font, PatternFill
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        print("owner-sheet: openpyxl is required; install it with: python3 -m pip install openpyxl", file=sys.stderr)
        return 2
    try:
        path = arguments.config.expanduser().resolve()
        if arguments.command == "init":
            init(path, arguments.language, arguments.title)
        else:
            config = load_config(path)
            if arguments.command == "build":
                build(config)
            else:
                read_or_ack(config, arguments.command, arguments.ids)
    except (ValueError, OSError, BadZipFile, InvalidFileException) as error:
        print("owner-sheet: " + str(error), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
